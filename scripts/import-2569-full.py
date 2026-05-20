"""
Parse Vet_to_be_2569_FINAL.xlsx รายชื่อ-แผนก sheet → emit SQL INSERT.

This is the canonical per-student list with:
  - student demographics (year, full name, nickname, รหัสนิสิต, phone, line/ig)
  - slot 1 (single dept + week + rank)
  - slot 2+ (multi-line, can be 1-3 lines for students with 2-4 total slots)

Output: full-seed-2569.sql ready to apply via Supabase apply_migration.
"""
import openpyxl
import re
import sys
from pathlib import Path

XLSX = Path("C:/Users/palmz/Desktop/Vet_to_be_2569_FINAL/Vet_to_be_2569_FINAL.xlsx")

# Thai dept name (canonical or with annotation) → slug
DEPT_BY_NAME = {
    "คลินิกโรคไต หัวใจ และระบบขับถ่ายปัสสาวะ": "nephrology-cardiac",
    "คลินิกโรคผิวหนัง": "dermatology",
    "คลินิกโรคมะเร็ง": "oncology",
    "คลินิกสัตว์พิเศษ": "exotic-pet",
    "แผนกชันสูตรโรคสัตว์ (ผ่าซาก)": "necropsy",
    "แผนกชันสูตรโรคสัตว์": "necropsy",
    "แผนกศัลยกรรม": "surgery",
    "แผนกสูติกรรม": "obstetrics",
    "แผนกอายุรกรรม": "internal-medicine",
    "หน่วยจักษุ": "ophthalmology",
    "หน่วยเวชระเบียน": "medical-records",
    "หน่วยเวชศาสตร์ฉุกเฉิน": "emergency",
    "หน่วยเวชศาสตร์ทางเลือก": "alternative-medicine",
    "หน่วยเวชศาสตร์ฟื้นฟู": "rehabilitation",
    "ห้องผ่าตัด": "operating-room",
    "ห้องยา": "pharmacy",
    "หออภิบาลสัตว์ป่วยขั้นวิกฤติ CCU": "ccu",
    "สำนักงาน (ประชาสัมพันธ์ PR VDO)": "pr-vdo",
    "สำนักงาน": "pr-vdo",
    "ศูนย์วิจัยโรคสัตว์น้ำ": "aquatic",
}

# Patterns
WEEK_RE = re.compile(r"^W(\d+)\s*\(([^)]*)\)\s*(?:\((\d+)\s*วัน\))?\s*(?:\(.*?\))?\s*$")
WEEK_RE_SIMPLE = re.compile(r"W(\d+)")
DAYS_RE = re.compile(r"\((\d+)\s*วัน\)")

# Rank tag → canonical
def parse_rank(tag):
    if not tag: return "manual"
    t = str(tag).strip()
    if "เพิ่มทีหลัง" in t or "แก้คนสลับ" in t:
        return "manual"
    if "อันดับ 1" in t: return "rank-1"
    if "อันดับ 2" in t: return "rank-2"
    if "อันดับ 3" in t: return "rank-3"
    if "เติม" in t: return "fill"
    if "สุ่ม" in t: return "random"
    return "manual"

def parse_source(tag):
    if not tag: return "algorithm"
    t = str(tag).strip()
    if "เพิ่มทีหลัง" in t: return "late-add"
    if "แก้คนสลับ" in t: return "manual-swap"
    return "algorithm"

def parse_week(week_str):
    """Return (week_num, days) from strings like 'W3 (28 พ.ค.-3 มิ.ย.) (4 วัน)'."""
    if not week_str:
        return None, 7
    s = str(week_str).strip()
    m = WEEK_RE_SIMPLE.search(s)
    if not m:
        return None, 7
    week = int(m.group(1))
    days_match = DAYS_RE.search(s)
    days = int(days_match.group(1)) if days_match else 7
    return week, days

def parse_dept_slug(dept_str):
    if not dept_str:
        return None
    v = str(dept_str).strip()
    # exact match first, then longest-prefix match
    if v in DEPT_BY_NAME:
        return DEPT_BY_NAME[v]
    best = None
    best_len = 0
    for name, slug in DEPT_BY_NAME.items():
        if v.startswith(name) and len(name) > best_len:
            best = slug
            best_len = len(name)
    return best

def parse_nickname(nick_full):
    """'ใบบัว#56' → ('ใบบัว', '056'). Handles 'ใบบัว #56', '#NN', etc."""
    if not nick_full:
        return None, None
    m = re.match(r"^(.+?)\s*#\s*0*(\d+)\s*$", str(nick_full).strip())
    if m:
        return m.group(1).strip(), m.group(2).zfill(3)
    return str(nick_full).strip(), None

def main():
    wb = openpyxl.load_workbook(XLSX, data_only=True)
    ws = wb["รายชื่อ-แผนก"]

    assignments = []
    students_seen = set()

    for r in range(5, ws.max_row + 1):
        seq = ws.cell(row=r, column=1).value
        if not seq:
            continue

        year_str = str(ws.cell(row=r, column=2).value or "")
        student_year = 5 if "5" in year_str else (4 if "4" in year_str else 4)
        full_name = str(ws.cell(row=r, column=3).value or "").strip()
        nick_raw = ws.cell(row=r, column=4).value
        full_id = str(ws.cell(row=r, column=5).value or "").strip()

        nickname, short_id_from_nick = parse_nickname(nick_raw)
        # short_id: prefer the # number in nickname, fallback to last 3 of รหัสนิสิต before "31" suffix
        if not short_id_from_nick and full_id:
            # Thai student IDs end in '31' (faculty code) — short id is digits 5..8
            short_id_from_nick = full_id[5:8] if len(full_id) >= 8 else full_id[-3:]
        short_id = short_id_from_nick or "?"

        students_seen.add(short_id)

        # Slot 1
        s1_dept = ws.cell(row=r, column=6).value
        s1_week = ws.cell(row=r, column=7).value
        s1_rank = ws.cell(row=r, column=8).value
        slug1 = parse_dept_slug(s1_dept)
        week1, days1 = parse_week(s1_week)
        if slug1 and week1:
            assignments.append({
                "dept_slug": slug1,
                "week": week1,
                "short_id": short_id,
                "nickname": nickname or "",
                "full_name": full_name,
                "student_year": student_year,
                "days_practiced": days1,
                "rank": parse_rank(s1_rank),
                "source": parse_source(s1_rank),
                "notes": None,
            })

        # Slot 2+ multi-line
        s2_depts_raw = ws.cell(row=r, column=9).value
        s2_weeks_raw = ws.cell(row=r, column=10).value
        s2_ranks_raw = ws.cell(row=r, column=11).value
        if s2_depts_raw:
            depts = [l.strip() for l in str(s2_depts_raw).split("\n") if l.strip()]
            weeks = [l.strip() for l in str(s2_weeks_raw or "").split("\n") if l.strip()]
            ranks = [l.strip() for l in str(s2_ranks_raw or "").split("\n") if l.strip()]
            for i, dept in enumerate(depts):
                slug = parse_dept_slug(dept)
                wk_str = weeks[i] if i < len(weeks) else None
                rk_str = ranks[i] if i < len(ranks) else None
                week, days = parse_week(wk_str)
                if slug and week:
                    assignments.append({
                        "dept_slug": slug,
                        "week": week,
                        "short_id": short_id,
                        "nickname": nickname or "",
                        "full_name": full_name,
                        "student_year": student_year,
                        "days_practiced": days,
                        "rank": parse_rank(rk_str),
                        "source": parse_source(rk_str),
                        "notes": None,
                    })

    print(f"-- {len(assignments)} assignments · {len(students_seen)} unique students", file=sys.stderr)

    # Audit · dept counts
    from collections import Counter
    dept_counts = Counter(a["dept_slug"] for a in assignments)
    print("-- per-dept:", file=sys.stderr)
    for slug, n in sorted(dept_counts.items(), key=lambda x: -x[1]):
        print(f"   {slug:24s} {n}", file=sys.stderr)

    # Emit SQL
    lines = []
    lines.append("-- Phase 2 full seed · regenerated from Vet_to_be_2569_FINAL.xlsx 'รายชื่อ-แผนก' sheet")
    lines.append("DELETE FROM public.vettobe_assignments WHERE year_id = 2569;")
    lines.append("")
    lines.append("INSERT INTO public.vettobe_assignments (year_id, dept_slug, week, short_id, nickname, full_name, student_year, days_practiced, rank, source, notes) VALUES")
    rows = []
    for a in assignments:
        sql_nick = (a["nickname"] or "").replace("'", "''")
        sql_name = (a["full_name"] or "").replace("'", "''")
        notes_sql = "NULL" if a["notes"] is None else f"'{str(a['notes']).replace(chr(39), chr(39)+chr(39))}'"
        rows.append(
            f"(2569, '{a['dept_slug']}', {a['week']}, '{a['short_id']}', "
            f"'{sql_nick}', '{sql_name}', {a['student_year']}, "
            f"{a['days_practiced']}, '{a['rank']}', '{a['source']}', {notes_sql})"
        )
    lines.append(",\n".join(rows) + ";")
    lines.append("")
    # Refresh year stats
    lines.append(f"UPDATE public.vettobe_years SET total_slots = {len(assignments)}, unique_students = {len(students_seen)}, updated_at = now() WHERE id = 2569;")

    out = Path(__file__).parent / "full-seed-2569.sql"
    out.write_text("\n".join(lines), encoding="utf-8")
    print(f"wrote {out} ({out.stat().st_size:,} bytes)", file=sys.stderr)

if __name__ == "__main__":
    main()
